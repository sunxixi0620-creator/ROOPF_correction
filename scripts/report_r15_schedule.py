"""R15 independent numerical, matched-case and actual API execution audit."""
from pathlib import Path
import csv,gzip,json,hashlib,collections,sys
from decimal import Decimal
import numpy as np
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
import r15_common as common
ARMS=list(common.CONFIGS)


def read(p):return json.loads(Path(p).read_text())
def csvwrite(p,rows):
    if not rows:return
    with Path(p).open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)


def report(exp,verify=False):
    from roopf.policy_api_r10 import validate_response,wire_bytes
    from roopf.current_evidence_agent import request_context_digest,digest
    exp=Path(exp);summaries=[];function_rows=[];trajectory_rows=[];api_counts={a:collections.Counter() for a in ARMS};consumed=set();credit=Decimal(0);credit_rows=0
    journals={};api_details=[];function_api={} 
    for p in (exp/'api_journal').glob('*.json'):
        r=read(p);journals[p.stem]=r
        assert hashlib.sha256(wire_bytes(r['request'])).hexdigest()==p.stem
        payload=json.loads(r['request']['messages'][-1]['content'].split('INPUT:\n',1)[1])
        assert payload['execution_context_sha256']==digest(r['execution_context'])
        if r['status'] in ('ok','partial'):
            decoded=validate_response(r['response'],payload)
            assert all(decoded[k]==r[k] for k in decoded)
        arm=r['execution_context']['arm'];c=api_counts[arm];c['attempts']+=1;c[r['status']]+=1
        c['valid_rows']+=len(r.get('valid_trajectory_ids',[]));c['fallback_rows']+=len(r.get('row_errors',{}))
        u=r.get('response',{}).get('usage',{});c['prompt_tokens']+=u.get('prompt_tokens',0);c['completion_tokens']+=u.get('completion_tokens',0)
        if isinstance(u.get('credit'),(int,float)):credit+=Decimal(str(u['credit']));credit_rows+=1
        ctx=r['execution_context'];key=(arm,ctx['group'],ctx['fid']);fc=function_api.setdefault(key,collections.Counter())
        fc.update(attempts=1,valid_rows=len(r.get('valid_trajectory_ids',[])),fallback_rows=len(r.get('row_errors',{})),prompt_tokens=u.get('prompt_tokens',0),completion_tokens=u.get('completion_tokens',0));fc[r['status']]+=1
        api_details.append(dict(config_id=arm,group=ctx['group'],fid=ctx['fid'],batch=ctx['batch'],identity=p.stem,status=r['status'],
            valid_rows=len(r.get('valid_trajectory_ids',[])),fallback_rows=len(r.get('row_errors',{})),prompt_tokens=u.get('prompt_tokens'),completion_tokens=u.get('completion_tokens'),
            provider_credit_units=u.get('credit'),error=r.get('error'),automatic_retry=False))
    csvwrite(exp/'API_REQUEST_DETAILS.csv',api_details)
    csvwrite(exp/'API_BY_CONFIGURATION_AND_FUNCTION.csv',[dict(config_id=a,group=g,fid=f,**dict(v)) for (a,g,f),v in function_api.items()])
    valid_coverage={};receipts=0
    for group,count in [('BBOB',24),('CEC-style',6),('UAV',5)]:
        for fid in range(1,count+1):
            per={a:[] for a in ARMS};refs={'frozen_final':[],'frozen_original':[]};coordinates={a:collections.Counter() for a in ARMS}
            for batch in range(10):
                folder=exp/'results'/f'{group}_f{fid:02d}'/f'b{batch:02d}'
                if not (folder/'REFERENCES.npz').exists():continue
                with np.load(folder/'REFERENCES.npz') as data:
                    frozen=data['frozen_final_values'].min(1).astype(float);original=data['original_values'].min(1).astype(float)
                    refs['frozen_final'].append((batch,frozen));refs['frozen_original'].append((batch,original))
                for arm in ARMS:
                    receipt=folder/(arm+'.json')
                    if not receipt.exists():continue
                    receipts+=1;r=read(receipt)
                    with np.load(folder/(arm+'.npz')) as data:values=data['values'].min(1).astype(float);points=data['executed_points'].copy()
                    per[arm].append((batch,values));coordinates[arm].update({k:r[k] for k in ['coordinate_changes_vs_frozen_final','coordinate_changes_vs_original','native_fe','cached_fe','actual_online_evaluated_points']})
                    if verify:
                        with gzip.open(folder/(arm+'_DECISIONS.json.gz'),'rt') as f:decisions=json.load(f)
                        assert len(decisions)==45
                        requests=[d for d in decisions if d['request_event']];assert [d['step'] for d in requests]==[common.CONFIGS[arm]['first_call_active_step'],common.CONFIGS[arm]['second_call_active_step']]
                        for d in decisions:
                            for i,c in enumerate(d['current']):
                                assert np.array_equal(points[i,55+d['step'],0],c['protected_point'])
                                assert np.array_equal(points[i,55+d['step'],1],c['candidates'][c['selected_index']])
                        for d in requests:
                            identity=d['api_request_sha256'];assert identity not in consumed
                            e=journals[identity];assert e['execution_context']==d['execution_context'] and e['config_ids']==d['configs']
                            assert e['decision_context_sha256']==request_context_digest(d['api_trajectories'],d['execution_context'])
                            assert d['api_row_validation']['valid_trajectory_ids']==e['valid_trajectory_ids']
                            for cont in decisions[d['step']+1:d['step']+common.CONFIGS[arm]['fusion_rounds']]:
                                assert cont['configs']==e['config_ids'] and cont['applied_generation']==[n.split('@')[0] for n in e['config_ids']]
                            consumed.add(identity);key=(group,fid,arm);valid_coverage[key]=valid_coverage.get(key,0)+len(e['valid_trajectory_ids'])
            complete={a:np.concatenate([v for b,v in sorted(rows)]) for a,rows in per.items() if len(rows)==10}
            complete_refs={a:np.concatenate([v for b,v in sorted(rows)]) for a,rows in refs.items() if len(rows)==10}
            for arm,values in complete.items():
                for reference,base in [*complete_refs.items(),*([('matched_default',complete['S02_L05'])] if 'S02_L05' in complete else [])]:
                    row=dict(group=group,fid=fid,arm=arm,first_call_fe=common.CONFIGS[arm]['first_call_fe'],fusion_rounds=common.CONFIGS[arm]['fusion_rounds'],reference=reference,n=100,
                        mean=float(values.mean()),population_std=float(values.std(ddof=0)),variance=float(values.var(ddof=0)),reference_mean=float(base.mean()),
                        mean_difference=float(values.mean()-base.mean()),wins=int((values<base).sum()),ties=int((values==base).sum()),losses=int((values>base).sum()),
                        primary=not(group=='BBOB' and fid==20),**dict(coordinates[arm]))
                    function_rows.append(row)
                    for i,(x,y) in enumerate(zip(values,base)):trajectory_rows.append(dict(group=group,fid=fid,arm=arm,reference=reference,trajectory=i,value=float(x),reference_value=float(y),difference=float(x-y)))
    for arm in ARMS:
        for group in ['BBOB','BBOB20_separate','CEC-style','UAV','primary_34_functions','BBOB24_supplementary']:
            for reference in ['frozen_final','matched_default','frozen_original']:
                rows=[r for r in function_rows if r['arm']==arm and r['reference']==reference and
                    ((group=='BBOB24_supplementary' and r['group']=='BBOB') or (group=='primary_34_functions' and r['primary']) or (group=='BBOB20_separate' and r['group']=='BBOB' and r['fid']==20) or
                     (group==r['group'] and not(group=='BBOB' and r['fid']==20)))]
                if rows:summaries.append(dict(arm=arm,first_call_fe=common.CONFIGS[arm]['first_call_fe'],fusion_rounds=common.CONFIGS[arm]['fusion_rounds'],group=group,reference=reference,functions=len(rows),
                    function_wins=sum(r['mean_difference']<0 for r in rows),function_ties=sum(r['mean_difference']==0 for r in rows),function_losses=sum(r['mean_difference']>0 for r in rows),
                    trajectory_wins=sum(r['wins'] for r in rows),trajectory_ties=sum(r['ties'] for r in rows),trajectory_losses=sum(r['losses'] for r in rows)))
    csvwrite(exp/'FUNCTION_STATISTICS_AND_COMPARISONS.csv',function_rows);csvwrite(exp/'PAIRED_TRAJECTORIES.csv',trajectory_rows);csvwrite(exp/'SUMMARY.csv',summaries)
    rankings=[]
    groups=['BBOB','CEC-style','UAV','primary_34_functions','BBOB20_separate','BBOB24_supplementary']
    for group in groups:
        rank_sums=collections.Counter();rank_counts=collections.Counter();best_counts=collections.Counter()
        keys={(r['group'],r['fid']) for r in function_rows if r['reference']=='frozen_final' and
            ((group=='BBOB24_supplementary' and r['group']=='BBOB') or (group=='primary_34_functions' and r['primary']) or (group=='BBOB20_separate' and r['group']=='BBOB' and r['fid']==20) or
             (group==r['group'] and not(group=='BBOB' and r['fid']==20)))}
        for g,f in keys:
            rows=[r for r in function_rows if r['reference']=='frozen_final' and r['group']==g and r['fid']==f]
            if len(rows)!=24:continue
            for r in rows:
                less=sum(v['mean']<r['mean'] for v in rows);equal=sum(v['mean']==r['mean'] for v in rows)
                rank_sums[r['arm']]+=less+(equal+1)/2;rank_counts[r['arm']]+=1
                best_counts[r['arm']]+=int(less==0)
        for arm in ARMS:
            if rank_counts[arm]:rankings.append(dict(group=group,arm=arm,first_call_fe=common.CONFIGS[arm]['first_call_fe'],fusion_rounds=common.CONFIGS[arm]['fusion_rounds'],functions=rank_counts[arm],
                average_rank=rank_sums[arm]/rank_counts[arm],best_or_tied_best_functions=best_counts[arm]))
    csvwrite(exp/'PARAMETER_RANKINGS.csv',rankings)
    execution_complete=receipts==8400
    if verify and execution_complete:assert consumed==set(journals) and len(journals)==16800
    valid_all=verify and all(valid_coverage.get((g,f,a),0)>0 for g,n in [('BBOB',24),('CEC-style',6),('UAV',5)] for f in range(1,n+1) for a in ARMS)
    if verify:
        common.validate_frozen()
    status='complete_live_schedule_measurement_verified' if execution_complete and valid_all else ('execution_complete_api_validation_incomplete' if execution_complete else 'partial')
    case_audit=audit_cases(exp) if verify else {}
    extra=export_views(exp,function_rows,rankings,summaries,api_counts)
    baseline_validation=dict(native_fe=0,cached_fe=0,batches=0)
    for path in exp.glob('results/*/b??/REFERENCE_VALIDATION.json'):
        c=read(path);baseline_validation['native_fe']+=c['native_fe'];baseline_validation['cached_fe']+=c['cached_fe'];baseline_validation['batches']+=1
    result=dict(baseline_validation=baseline_validation,case_audit=case_audit,views=extra,status=status,completed_arm_batches=receipts,expected_arm_batches=8400,matched_actual_api_records=len(consumed),expected_actual_requests=16800,
        actual_api={k:dict(v) for k,v in api_counts.items()},provider_credit_units=str(credit),responses_with_credit=credit_rows,currency_inferred=False,
        summaries=summaries,parameter_rankings=rankings,not_significance_test=True,bbob20_separate=True)
    write=exp/('FINAL_RESULT.json' if execution_complete else 'PARTIAL_RESULT.json');write.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    return result

def audit_cases(exp):
    from roopf.scheduled_evidence_agent import ScheduledEvidenceAgent,digest
    stats=[]
    for cid,config in common.CONFIGS.items():
        folder=exp/'cases'/cid
        if not (folder/'offline_memory.json').exists():continue
        memory=read(folder/'offline_memory.json')
        assert memory['schedule']==common.schedule(cid) and memory['schedule_signature']==digest(common.schedule(cid))
        assert len(memory['cases'])==464
        with gzip.open(folder/'HELDOUT_CASES.json.gz','rt') as f:held=json.load(f)
        assert len(held)==112
        allcases=memory['cases']+held
        assert all(r['schedule_signature']==memory['schedule_signature'] for r in allcases)
        assert {r['step'] for r in allcases}=={config['first_call_active_step'],config['second_call_active_step']}
        assert all(r['nfe_before']==210+2*r['step'] for r in allcases)
        summary=read(folder/'COLLECTION_SUMMARY.json');assert summary['native_fe']<=3000000
        stats.append(summary)
    if len(stats)==24:
        # Rebuilt default must equal frozen case semantics; only new identity metadata may differ.
        default=read(exp/'cases/S02_L05/offline_memory.json')
        old=read(common.ROOT/'experiments/r13_cec_uniform/offline_memory.json')
        assert default['coverage_radii']==old['coverage_radii'],'Default coverage changed'
        for x,y in zip(default['cases'],old['cases']):
            for key in ['id','function','features','advantages','prefix_profile','reference_profile','step','initial_radius']:
                assert x[key]==y[key],('Default case semantic mismatch',key,x['id'])
    return dict(configurations=len(stats),contexts=576*len(stats),native_fe=sum(s['native_fe'] for s in stats),
                cached_fe=sum(s['cached_fe'] for s in stats),summaries=stats,real_api_calls=0)

def export_views(exp,rows,ranks,summaries,api):
    views=[]
    for r in ranks:
        c=common.CONFIGS[r['arm']]
        views.append(dict(r,view='factorial_6x4',timing_group=c['timing_group']))
    for view,subset in [('timing_at_L05',[k for k,c in common.CONFIGS.items() if c['fusion_rounds']==5]),
                        ('duration_at_S02',[k for k,c in common.CONFIGS.items() if c['timing_group']=='S02'])]:
        for group in ['BBOB','CEC-style','UAV','primary_34_functions','BBOB20_separate','BBOB24_supplementary']:
            totals=collections.Counter();counts=collections.Counter()
            chosen=[r for r in rows if r['arm'] in subset and r['reference']=='frozen_final' and
                ((group=='primary_34_functions' and r['primary']) or (group=='BBOB24_supplementary' and r['group']=='BBOB') or
                 (group=='BBOB20_separate' and r['group']=='BBOB' and r['fid']==20) or
                 (group==r['group'] and not (group=='BBOB' and r['fid']==20)))]
            for g,f in {(r['group'],r['fid']) for r in chosen}:
                rr=[r for r in chosen if r['group']==g and r['fid']==f]
                if len(rr)!=len(subset):continue
                for r in rr:
                    totals[r['arm']]+=sum(x['mean']<r['mean'] for x in rr)+(sum(x['mean']==r['mean'] for x in rr)+1)/2
                    counts[r['arm']]+=1
            for cid in subset:
                if counts[cid]:views.append(dict(group=group,arm=cid,first_call_fe=common.CONFIGS[cid]['first_call_fe'],fusion_rounds=common.CONFIGS[cid]['fusion_rounds'],functions=counts[cid],
                    average_rank=totals[cid]/counts[cid],best_or_tied_best_functions=None,view=view,timing_group=common.CONFIGS[cid]['timing_group']))
    csvwrite(exp/'FACTORIAL_AND_ONE_FACTOR_VIEWS.csv',views)
    # Coverage and effective control are reported from the actual API input/execution, including failures.
    coverage=[]
    for cid in ARMS:
        for g,n in [('BBOB',24),('CEC-style',6),('UAV',5)]:
            for fid in range(1,n+1):
                c=collections.Counter()
                for batch in range(10):
                    path=exp/'results'/f'{g}_f{fid:02d}'/f'b{batch:02d}'/(cid+'_DECISIONS.json.gz')
                    if not path.exists():continue
                    with gzip.open(path,'rt') as f:decisions=json.load(f)
                    for d in decisions:
                        c['nondefault_action_trajectory_rounds']+=sum(x not in ('matched38','release') for x in d['configs'])
                        if d['request_event']:
                            c['query_rows']+=len(d['api_trajectories'])
                            c['nondefault_requested_rows']+=sum(x!='matched38' for x in d['configs'])
                            for t in d['api_trajectories']:
                                c['covered_rows']+=int(t['evidence']['supported'])
                                c['default_only_allowed_rows']+=int(t['eligible_configs']==['matched38'])
                if c:
                    coverage.append(dict(config_id=cid,group=g,fid=fid,**dict(c),default_only_fraction=c['default_only_allowed_rows']/max(1,c['query_rows'])))
    csvwrite(exp/'CASE_COVERAGE_AND_EFFECTIVE_CONTROL.csv',coverage)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font,PatternFill,Alignment
        from openpyxl.utils import get_column_letter
        wb=Workbook();wb.remove(wb.active)
        for name,data in [('参数排名',views),('分类胜平负',summaries),('各函数测量值',rows),('案例覆盖和实际控制',coverage),('API汇总',[dict(config_id=k,**dict(v)) for k,v in api.items()])]:
            ws=wb.create_sheet(name)
            if not data:continue
            keys=list(dict.fromkeys(k for r in data for k in r));ws.append(keys)
            for r in data:ws.append([r.get(k) for k in keys])
            ws.freeze_panes='D2';ws.auto_filter.ref=ws.dimensions
            for cell in ws[1]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='193B60');cell.alignment=Alignment(wrap_text=True,vertical='center')
            ws.row_dimensions[1].height=36
            for j,k in enumerate(keys,1):ws.column_dimensions[get_column_letter(j)].width=min(40,max(16,len(k)+2))
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    if cell.row%2==0:cell.fill=PatternFill('solid',fgColor='EDF3F9')
                    if isinstance(cell.value,float):cell.number_format='0.000000E+00'
        wb.save(exp/'LLM_Agent_匹配时序参数实验结果.xlsx')
    except ImportError:pass
    return dict(rows=len(views),coverage_rows=len(coverage))

if __name__=='__main__':
    import argparse
    p=argparse.ArgumentParser();p.add_argument('--verify',action='store_true');a=p.parse_args()
    print(json.dumps(report(common.EXP,a.verify)))
