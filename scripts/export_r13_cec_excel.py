"""Produce a new workbook with verified R13 CEC rows; preserve the frozen XLSX."""
from pathlib import Path
from copy import copy
from collections import Counter
import csv
import hashlib
import json
import re
from openpyxl import load_workbook
from openpyxl.styles import Font,PatternFill,Alignment

ROOT=Path(__file__).resolve().parents[1]
EXP=ROOT/'experiments/r13_cec_uniform'
SOURCE=ROOT/'exports/当前Full_论文对照与消融_BBOB24_按算法分类_说明置顶.xlsx'
TARGET=ROOT/'exports/当前Full_论文对照与消融_BBOB24_CEC统一流程重测_R13.xlsx'


def main():
    frozen_hash=hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    assert frozen_hash=='f6841509e51126f3ab693313a4184fb2f876d33e4feb2c4d663967ea4c8b1ff5'
    result=json.loads((EXP/'FINAL_RESULT.json').read_text())
    assert result['status']=='complete_live_cec_uniform_measurement_verified'
    with (EXP/'FUNCTION_STATISTICS.csv').open(encoding='utf-8-sig',newline='') as f:
        rows=list(csv.DictReader(f))
    stats={(r['arm'],int(r['fid'])):r for r in rows}
    book=load_workbook(SOURCE)
    unchanged={s.title:tuple(tuple(c.value for c in row) for row in s) for s in book if 'BBOB' in s.title or 'UAV' in s.title}
    detail=book['均值方差_论文对照_CEC-style']
    for row in detail.iter_rows(min_row=4):
        label,fid=row[0].value,int(re.search(r'(\d+)$',row[1].value).group(1))
        if label=='当前 Full':
            r=stats['llm_full',fid];row[2].value=float(r['mean']);row[3].value=float(r['variance'])
        elif label=='ROOPF 原版':
            r=stats['frozen_baseline',fid]
            assert abs(row[2].value-float(r['mean']))<=1e-14*max(abs(float(r['mean'])),1e-30)
            assert abs(row[3].value-float(r['variance']))<=1e-14*max(abs(float(r['variance'])),1e-30)
            row[2].value=float(r['mean']);row[3].value=float(r['variance'])
    values={}
    for row in detail.iter_rows(min_row=4):
        values.setdefault(row[0].value,{})[int(re.search(r'(\d+)$',row[1].value).group(1))]=row[2].value
    function_ranks={name:[] for name in values}
    for fid in range(1,7):
        ordered=sorted(values,key=lambda a:values[a][fid]);start=0
        while start<len(ordered):
            end=start+1
            while end<len(ordered) and values[ordered[end]][fid]==values[ordered[start]][fid]:end+=1
            rank=(start+1+end)/2
            for name in ordered[start:end]:function_ranks[name].append(rank)
            start=end
    averages={a:sum(r)/6 for a,r in function_ranks.items()}
    ordered=sorted(averages,key=lambda a:(averages[a],0 if a=='当前 Full' else 1,a))
    sheet=book['对照_CEC-style'];styles={}
    for row in sheet.iter_rows(min_row=7,max_row=22):styles[row[1].value]=[copy(c._style) for c in row]
    counts=Counter(averages.values())
    for i,name in enumerate(ordered,7):
        avg=averages[name];rank=1+sum(v<avg for v in averages.values())
        full=[values['当前 Full'][f] for f in range(1,7)];other=[values[name][f] for f in range(1,7)]
        w,t,l=(sum(a<b for a,b in zip(full,other)),sum(a==b for a,b in zip(full,other)),sum(a>b for a,b in zip(full,other)))
        data=[('并列' if counts[avg]>1 else '')+str(rank),name,avg,None if name=='当前 Full' else w,None if name=='当前 Full' else t,None if name=='当前 Full' else l,6]
        for j,v in enumerate(data,1):
            c=sheet.cell(i,j);c.value=v;c._style=copy(styles[name][j-1])
            if name=='当前 Full':c.fill=PatternFill('solid',fgColor='E4F1E9');c.font=Font(name=c.font.name,size=c.font.sz,bold=True,color='193B60')
    sheet['A2']='排名按每函数最终均值从低到高计算，再取平均秩；当前Full使用CEC统一流程重测结果，与冻结ROOPF基线同时参排。\nCEC六函数均开放与BBOB一致的生成、Residual、案例准入、LLM及保护选择路径；API有效返回与实际执行已核对。\n其余算法采用论文归档均值，未在本轮重跑，归档偏移一致性仍未确认；排名仅供参考，不代表配对显著性。'
    sheet['A2'].alignment=Alignment(wrap_text=True,vertical='center')
    ablation=book['消融_CEC-style']
    mapping={'llm_no_anchor':8,'llm_no_online_generators':9,'no_agent':10}
    for r in result['comparisons']:
        if r['reference'] not in mapping:continue
        row=mapping[r['reference']]
        for col,key in enumerate(['function_wins','function_ties','function_losses','trajectory_wins','trajectory_ties','trajectory_losses'],2):ablation.cell(row,col,r[key])
    ablation['A2']='移除Anchor：移除预训练Anchor网络及其候选生成路径，以原生信赖域候选补足两个参考位置；其余模块与Full冻结案例保留，Anchor温度控制失效。\n移除在线部分：完全移除六个在线生成算子及额外原生提议，保留Anchor、代理、Residual及Agent。\n移除Agent：移除LLM、检索准入、Agent历史、额外提议、参数控制和重排序；保留Anchor、六算子、代理及Residual。\n本表四个条件均采用CEC统一执行策略，保持相同初始种群与目标评价随机流。'
    ablation['A2'].alignment=Alignment(wrap_text=True,vertical='center')
    variants={'当前 Full':'llm_full','移除 Anchor':'llm_no_anchor','移除在线部分':'llm_no_online_generators','移除 Agent':'no_agent'}
    sheet=book['均值方差_消融_CEC-style']
    for row in sheet.iter_rows(min_row=4):
        fid=int(re.search(r'(\d+)$',row[1].value).group(1));r=stats[variants[row[0].value],fid]
        row[2].value=float(r['mean']);row[3].value=float(r['variance'])
    notes=book['说明与来源']
    changes={4:'BBOB/UAV沿用冻结Full及消融证据；CEC使用R13统一流程下的Full、冻结基线与三项整模块消融重测。',
        10:'CEC六函数已解除算法中的专属结构排除，采用与BBOB一致的主流程和Agent准入；目标函数及边界保留。论文归档偏移一致性仍未确认。',
        12:'BBOB/UAV沿用冻结的严格在线生成模块消融。CEC使用R13结果；仍冻结Full案例与门限，局部尺度动作在该消融中为失效别名，并向LLM披露。',
        14:'移除Agent与Full使用相同CUDA代理；冻结ROOPF基线保持CPU路径。CEC的两者均完整重测并报告，计算数值差异包含在比较中。',
        15:'均D10、每函数100条轨迹、FE300含初始100；BBOB24/UAV5保留既有结果，CEC6重测5个条件共3000条轨迹，使用真实LLM并保持配对初始种群与201个目标评价随机流边界。',
        16:'BBOB/UAV：experiments/r10_current_candidate_fusion/FUNCTION_COMPARISON.csv；CEC：experiments/r13_cec_uniform/FUNCTION_STATISTICS.csv',
        18:'BBOB/UAV：experiments/r12_strict_component_ablations/WHOLE_MODULE_ABLATION_SUMMARY.csv；CEC：experiments/r13_cec_uniform/SUMMARY.csv',
        19:'BBOB/UAV：experiments/r12_strict_component_ablations/WHOLE_MODULE_FUNCTION_COMPARISONS.csv；CEC：experiments/r13_cec_uniform/FUNCTION_COMPARISONS.csv',
        20:'冻结Excel未覆盖；BBOB/UAV工作表逐单元格值核对不变。CEC替换为已完成重测数据，重新计算均值、总体方差、平均秩和胜平负；完整API映射见R13 API_EXECUTION_AUDIT.json。'}
    for row,text in changes.items():notes.cell(row,2,text);notes.cell(row,2).alignment=Alignment(wrap_text=True,vertical='center');notes.row_dimensions[row].height=max(notes.row_dimensions[row].height or 20,60)
    assert all(tuple(tuple(c.value for c in row) for row in book[name])==v for name,v in unchanged.items())
    book.save(TARGET)
    checked=load_workbook(TARGET,data_only=False)
    assert all(tuple(tuple(c.value for c in row) for row in checked[name])==v for name,v in unchanged.items())
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==frozen_hash
    record=dict(file=str(TARGET.relative_to(ROOT)),sha256=hashlib.sha256(TARGET.read_bytes()).hexdigest(),
        frozen_workbook_sha256=frozen_hash,untouched_BBOB_UAV_sheets=list(unchanged),cec_function_ranks=averages,
        cec_final_means_variances_updated=True,cec_ablation_counts_updated=True,no_cross_category_total=True,
        measurement_columns=['算法/版本','函数','均值','方差'])
    (EXP/'EXCEL_EXPORT.json').write_text(json.dumps(record,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(record,ensure_ascii=False))


if __name__=='__main__':main()
