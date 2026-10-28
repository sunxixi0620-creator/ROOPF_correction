"""Report paired FE200/FE210 results and preserve the older admission sweep."""
import csv,json,hashlib
from pathlib import Path
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
import r16_common as c

def readcsv(p):
    with p.open(encoding='utf-8-sig') as f:return list(csv.DictReader(f))
def writecsv(p,rows):
    with p.open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
def main():
    assert c.read(c.EXP/'FINAL_RESULT.json')['status']=='complete_live_schedule_measurement_verified'
    assert c.read(c.EXP/'INDEPENDENT_FINAL_AUDIT.json')['status']=='complete_live_schedule_measurement_independently_verified'
    roots={'R14':c.ROOT/'experiments/r14_online_entry_timing','R15':c.ROOT/'experiments/r15_llm_schedule_matched','R16':c.EXP}
    measurements=[];hashes={};paired=[]
    for name,root in roots.items():
        p=root/'FUNCTION_STATISTICS_AND_COMPARISONS.csv';hashes[str(p.relative_to(c.ROOT))]=c.sha(p)
        for r in readcsv(p):
            if r['reference']!='frozen_final':continue
            measurements.append(dict(round=name,config=r['arm'],online_fe=int(r['threshold']) if name=='R14' else (210 if name=='R15' else 200),group=r['group'],function=int(r['fid']),mean=float(r['mean']),population_std=float(r['population_std']),variance=float(r['variance'])))
    for group,n in [('BBOB',24),('CEC-style',6),('UAV',5)]:
        for fid in range(1,n+1):
            for cid in c.CONFIGS:
                vals={k:[] for k in ('R15','R16')}
                for b in range(10):
                    suffix=Path('results')/f'{group}_f{fid:02d}'/f'b{b:02d}'
                    with np.load(roots['R15']/suffix/'REFERENCES.npz') as a,np.load(roots['R16']/suffix/'REFERENCES.npz') as z:
                        for k in a.files:assert np.array_equal(a[k],z[k]),(group,fid,b,k)
                    for name in vals:
                        with np.load(roots[name]/suffix/(cid+'.npz')) as a:vals[name].extend(a['values'].min(1).astype(float))
                x,y=np.array(vals['R16']),np.array(vals['R15'])
                paired.append(dict(config=cid,group=group,function=fid,mean_fe200=float(x.mean()),mean_fe210=float(y.mean()),difference=float(x.mean()-y.mean()),wins=int((x<y).sum()),ties=int((x==y).sum()),losses=int((x>y).sum())))
    groups={'BBOB24':lambda g,f:g=='BBOB','BBOB23':lambda g,f:g=='BBOB' and f!=20,'BBOB20':lambda g,f:g=='BBOB' and f==20,'CEC6':lambda g,f:g=='CEC-style','UAV5':lambda g,f:g=='UAV','primary34':lambda g,f:(g,f)!=('BBOB',20),'all35':lambda g,f:True}
    rankings=[];summaries=[]
    for scope in [('R14',),('R15',),('R16',),('R15','R16')]:
        rows=[r for r in measurements if r['round'] in scope]
        funcs={(r['group'],r['function']) for r in rows};totals={}
        for g,f in funcs:
            rr=[r for r in rows if (r['group'],r['function'])==(g,f)]
            for r in rr:
                rank=sum(q['mean']<r['mean'] for q in rr)+(sum(q['mean']==r['mean'] for q in rr)+1)/2
                for label,pred in groups.items():
                    if pred(g,f):totals.setdefault((label,r['round'],r['config']),[]).append(rank)
        for (label,rd,cid),ranks in totals.items():rankings.append(dict(comparison='+'.join(scope),group=label,round=rd,config=cid,functions=len(ranks),average_rank=float(np.mean(ranks))))
    for cid in c.CONFIGS:
        for label,pred in groups.items():
            rr=[r for r in paired if r['config']==cid and pred(r['group'],r['function'])]
            summaries.append(dict(config=cid,group=label,function_wins=sum(r['difference']<0 for r in rr),function_ties=sum(r['difference']==0 for r in rr),function_losses=sum(r['difference']>0 for r in rr),trajectory_wins=sum(r['wins'] for r in rr),trajectory_ties=sum(r['ties'] for r in rr),trajectory_losses=sum(r['losses'] for r in rr)))
    tables={'各轮均值方差':measurements,'同配置200对210逐函数':paired,'同配置分类胜平负':summaries,'分类参数排名':rankings}
    wb=Workbook();wb.remove(wb.active)
    for name,rows in tables.items():
        writecsv(c.EXP/('综合_'+name+'.csv'),rows)
        ws=wb.create_sheet(name);keys=list(rows[0]);ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(keys));ws.cell(1,1,name+'；目标越低越好，排名仅在同一 comparison 内比较，BBOB20单列主分析。')
        ws.append(keys)
        for r in rows:ws.append([r[k] for k in keys])
        ws.freeze_panes='D3';ws.auto_filter.ref=f'A2:{get_column_letter(len(keys))}{ws.max_row}'
        for row in ws:
            for cell in row:
                cell.alignment=Alignment(vertical='center',wrap_text=True)
                if cell.row<=2:cell.fill=PatternFill('solid',fgColor='193B60');cell.font=Font(color='FFFFFF',bold=True)
                elif cell.row%2==0:cell.fill=PatternFill('solid',fgColor='EDF3F9')
                if isinstance(cell.value,float):cell.number_format='0.000000E+00'
        for i in range(1,len(keys)+1):ws.column_dimensions[get_column_letter(i)].width=22
        ws.row_dimensions[1].height=44;ws.row_dimensions[2].height=35
    wb.save(c.EXP/'在线200FE与既有参数实验综合结果.xlsx')
    c.write(c.EXP/'CROSS_ROUND_SUMMARY_BINDING.json',dict(status='completed_paired_comparison',source_sha256=c.sha(__file__),inputs=hashes,reference_batches_checked=350,comparison_is_cross_run=True,matched_cases_rebuilt=True,api_nondeterminism_preserved=True))
if __name__=='__main__':main()
