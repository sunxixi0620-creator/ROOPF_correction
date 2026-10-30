"""Validate R19 evidence, pair trajectory values and export ablation results."""
import csv,json,gzip,hashlib,collections,time
from decimal import Decimal
from pathlib import Path
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
from roopf.policy_api_r10 import validate_response,wire_bytes
from roopf.current_evidence_agent import request_context_digest,digest
from experiment_r19_ablation import E,ARMS,binding,check_records

def csvout(name,rows):
 if not rows:return
 with (E/name).open('w',encoding='utf-8-sig',newline='') as f:
  w=csv.DictWriter(f,fieldnames=list(dict.fromkeys(k for x in rows for k in x)));w.writeheader();w.writerows(rows)
def report():
 identity=binding()
 assert len(list(E.glob('results/*/b??/BATCH_COMPLETE.json')))==350
 events={};api=collections.Counter();credit=Decimal(0);credit_count=0
 for p in (E/'api_journal').glob('*.json'):
  v=json.loads(p.read_text());assert hashlib.sha256(wire_bytes(v['request'])).hexdigest()==p.stem
  payload=json.loads(v['request']['messages'][-1]['content'].split('INPUT:\n',1)[1]);assert payload['execution_context_sha256']==digest(v['execution_context'])
  if v['status'] in ('ok','partial'):
   decoded=validate_response(v['response'],payload);assert all(decoded[k]==v[k] for k in decoded)
  assert v['status'] in ('ok','partial','failed');api[v['status']]+=1;api['attempts']+=1;api['valid_rows']+=len(v.get('valid_trajectory_ids',[]));api['fallback_rows']+=len(v.get('row_errors',{}))
  u=v.get('response',{}).get('usage',{});api['prompt_tokens']+=u.get('prompt_tokens',0);api['completion_tokens']+=u.get('completion_tokens',0)
  if isinstance(u.get('credit'),(int,float)):credit+=Decimal(str(u['credit']));credit_count+=1
  events[p.stem]=v
 assert len(events)==2100
 stats=[];pairs=[];consumed=set();cov=collections.Counter();native=collections.Counter();changed=collections.Counter()
 for g,n in [('BBOB',24),('CEC-style',6),('UAV',5)]:
  for fid in range(1,n+1):
   vals={a:[] for a in ARMS};refs={'original':[],'archived_full':[]};points={a:[] for a in ARMS}
   for b in range(10):
    folder=E/'results'/f'{g}_f{fid:02d}'/f'b{b:02d}'
    with np.load(folder/'REFERENCES.npz') as z:
     for name in refs:refs[name].extend(z[name+'_values'].min(1))
    for arm in ARMS:
     receipt=json.loads((folder/(arm+'.json')).read_text());assert receipt['binding']==identity and receipt['objective_rng_boundaries_matched']==201
     native[arm+'_native']+=receipt['native_fe'];native[arm+'_cached']+=receipt['cached_fe']
     with np.load(folder/(arm+'.npz')) as z:v=z['values'];pt=z['executed_points'];ref=z['reference_points'];assert v.shape==(10,300) and np.isfinite(v).all();vals[arm].extend(v.min(1));points[arm].append(pt.copy())
     with gzip.open(folder/(arm+'_DECISIONS.json.gz'),'rt') as f:rr=json.load(f)
     if arm!='no_agent':
      check_records(rr,pt,ref,arm)
      for rec in rr:
       if not rec['request_event']:continue
       identity=rec['api_request_sha256'];assert identity not in consumed;j=events[identity];assert j['execution_context']==rec['execution_context'] and j['config_ids']==rec['configs']
       assert j['decision_context_sha256']==request_context_digest(rec['api_trajectories'],rec['execution_context'])
       assert rec['api_row_validation']['valid_trajectory_ids']==j['valid_trajectory_ids'];consumed.add(identity);cov[(g,fid,arm)]+=len(j['valid_trajectory_ids'])
     else:assert not rr
   for a in vals:vals[a]=np.asarray(vals[a],dtype=float);points[a]=np.concatenate(points[a])
   for name in refs:refs[name]=np.asarray(refs[name],dtype=float)
   for arm in ARMS:
    v=vals[arm];stats.append(dict(group=g,fid=fid,arm=arm,n=100,mean=float(v.mean()),population_std=float(v.std(ddof=0)),variance=float(v.var(ddof=0))))
    changed[(g,fid,arm)]=int((points[arm]!=points['llm_full']).any(-1).sum())
   for other in ['llm_no_anchor','llm_no_online_generators','no_agent','original','archived_full']:
    x=vals['llm_full'];y=vals.get(other,refs.get(other));pairs.append(dict(group=g,fid=fid,comparison=other,full_mean=float(x.mean()),other_mean=float(y.mean()),mean_difference=float(x.mean()-y.mean()),wins=int((x<y).sum()),ties=int((x==y).sum()),losses=int((x>y).sum()),changed_points_vs_full=changed.get((g,fid,other))))
 assert consumed==set(events) and len(cov)==105 and all(v>0 for v in cov.values()),'Valid executed API coverage missing'
 summaries=[]
 for label,pred in [('BBOB23',lambda g,f:g=='BBOB' and f!=20),('BBOB20',lambda g,f:g=='BBOB' and f==20),('BBOB24',lambda g,f:g=='BBOB'),('CEC-style',lambda g,f:g=='CEC-style'),('UAV',lambda g,f:g=='UAV'),('primary34',lambda g,f:(g,f)!=('BBOB',20))]:
  for comp in ['llm_no_anchor','llm_no_online_generators','no_agent','original','archived_full']:
   rr=[r for r in pairs if r['comparison']==comp and pred(r['group'],r['fid'])]
   summaries.append(dict(group=label,comparison=comp,functions=len(rr),function_wins=sum(r['mean_difference']<0 for r in rr),function_ties=sum(r['mean_difference']==0 for r in rr),function_losses=sum(r['mean_difference']>0 for r in rr),trajectory_wins=sum(r['wins'] for r in rr),trajectory_ties=sum(r['ties'] for r in rr),trajectory_losses=sum(r['losses'] for r in rr),changed_points_vs_full=sum(r['changed_points_vs_full'] or 0 for r in rr) if comp in ARMS else None))
 csvout('FUNCTION_STATISTICS.csv',stats);csvout('PAIRED_COMPARISONS.csv',pairs);csvout('CATEGORY_SUMMARY.csv',summaries)
 result=dict(status='complete_live_ablation_verified',trajectories=14000,condition_batches=1400,api=dict(api),api_provider_credit_units=str(credit),credit_responses=credit_count,currency_inferred=False,coverage_pairs=len(cov),native_cached_fe=dict(native),bbob20_separate=True,cec_unlocked=True,summaries=summaries)
 wb=Workbook();wb.remove(wb.active)
 for title,rows in [('分类胜平负',summaries),('逐函数均值方差',stats),('逐函数配对对照',pairs)]:
  ws=wb.create_sheet(title);ks=list(rows[0]);ws.append(ks)
  for r in rows:ws.append([r.get(k) for k in ks])
  ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions
  for c in ws[1]:c.fill=PatternFill('solid',fgColor='193B60');c.font=Font(color='FFFFFF',bold=True);c.alignment=Alignment(wrap_text=True)
  for j in range(1,len(ks)+1):ws.column_dimensions[get_column_letter(j)].width=21
  for row in ws.iter_rows(min_row=2):
   for c in row:
    if c.row%2==0:c.fill=PatternFill('solid',fgColor='EDF3F9')
    if isinstance(c.value,float):c.number_format='0.000000E+00'
 wb.save(E/'S01L08_三项消融结果.xlsx')
 (E/'FINAL_RESULT.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n');(E/'EXECUTION_STATUS.json').write_text(json.dumps({'status':result['status'],'finished':time.strftime('%Y-%m-%dT%H:%M:%S%z')},indent=2)+'\n')
 print(json.dumps(dict(status=result['status'],trajectories=14000,api=dict(api))),flush=True)
if __name__=='__main__':report()
