"""Create a new results workbook from audited R19 and archived paper values."""
from pathlib import Path
from copy import copy
from collections import defaultdict
import csv,json,hashlib,math
import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
R=Path(__file__).resolve().parents[2];OUT=Path(__file__).resolve().parent
EXP=R/'experiments/r19_s01l08_component_ablations'
SRC=R/'exports/当前Full_论文对照与消融_BBOB24_CEC统一流程重测_R13.xlsx'
TARGET=OUT/'S01L08当前版本_论文对照与三项消融_BBOB24.xlsx'
assert json.loads((EXP/'INDEPENDENT_AUDIT.json').read_text())['status']=='complete_independently_verified'
assert json.loads((EXP/'FINAL_RESULT.json').read_text())['status']=='complete_live_ablation_verified'
def read(name):
 with (EXP/name).open(encoding='utf-8-sig') as f:return list(csv.DictReader(f))
stats={(x['group'],int(x['fid']),x['arm']):x for x in read('FUNCTION_STATISTICS.csv')}
summ={(x['group'],x['comparison']):x for x in read('CATEGORY_SUMMARY.csv')}
comparison={(x['group'],int(x['fid']),x['comparison']):x for x in read('PAIRED_COMPARISONS.csv')}
labels={'当前 Full':'llm_full','移除 Anchor':'llm_no_anchor','移除在线部分':'llm_no_online_generators','移除 Agent':'no_agent'}
w=load_workbook(SRC);validation={}
for group,n in [('BBOB',24),('CEC-style',6),('UAV',5)]:
 originals={}
 for fid in range(1,n+1):
  v=[]
  for b in range(10):
   with np.load(EXP/'results'/f'{group}_f{fid:02d}'/f'b{b:02d}'/'REFERENCES.npz') as a:v.extend(a['original_values'].min(1).astype(float))
  z=np.asarray(v,dtype=float);assert z.shape==(100,);originals[fid]=(float(z.mean()),float(z.var(ddof=0)))
  row=comparison[group,fid,'original'];assert float(row['other_mean'])==originals[fid][0]
 for suffix in ['论文对照','消融']:
  sheet=w[f'均值方差_{suffix}_{group}']
  for row in sheet.iter_rows(min_row=4):
   name,func=row[0].value,row[1].value
   if name is None:continue
   fid=int(str(func).replace(group,''));assert 1<=fid<=n
   if name in labels:
    d=stats[group,fid,labels[name]];row[2].value=float(d['mean']);row[3].value=float(d['variance'])
   elif name=='ROOPF 原版':row[2].value,row[3].value=originals[fid]
  sheet['A2']='按算法分类，数值为每个函数 100 条轨迹的最终均值与方差；论文归档方差沿用原表。'
  if 'A2:D2' not in [str(x) for x in sheet.merged_cells.ranges]:sheet.merge_cells('A2:D2')
  sheet['A2'].alignment=Alignment(wrap_text=True,vertical='center');sheet.row_dimensions[2].height=32
 detail=w[f'均值方差_论文对照_{group}'];data=defaultdict(dict)
 for name,func,mean,var in detail.iter_rows(min_row=4,values_only=True):
  if name is not None:data[name][int(str(func).replace(group,''))]=float(mean)
 assert all(set(d)==set(range(1,n+1)) for d in data.values())
 scores={k:0. for k in data}
 for fid in range(1,n+1):
  vals=[v[fid] for v in data.values()]
  for name,arr in data.items():
   value=arr[fid];scores[name]+=sum(x<value for x in vals)+(sum(x==value for x in vals)+1)/2
 scores={a:v/n for a,v in scores.items()};order=sorted(data,key=lambda a:(scores[a],a));sheet=w[f'对照_{group}']
 sheet['A2']='按每个函数最终均值排名，再取平均秩；胜平负以当前 Full 为参照。论文其他算法沿用归档数据，仅供参考；Full 与 ROOPF 基线来自本轮配对记录。'+ ('\n包含 BBOB1—24；BBOB20具有批次耦合随机语义，主分析口径需单独解释。' if group=='BBOB' else ('\nCEC-style 已解除结构锁；论文归档的任务偏移一致性仍未完全确认。' if group=='CEC-style' else ''))
 sheet.row_dimensions[2].height=58
 for j,name in enumerate(order,7):
  value=scores[name];pos=sum(scores[o]<value for o in order)+1;tied=sum(scores[o]==value for o in order)>1
  wr=sum(data['当前 Full'][f]<data[name][f] for f in range(1,n+1));tr=sum(data['当前 Full'][f]==data[name][f] for f in range(1,n+1));lr=n-wr-tr
  cells=[('并列'+str(pos)) if tied else str(pos),name,value,None if name=='当前 Full' else wr,None if name=='当前 Full' else tr,None if name=='当前 Full' else lr,n]
  for col,v in enumerate(cells,1):
   c=sheet.cell(j,col);c.value=v;c.alignment=Alignment(horizontal='center' if col!=2 else 'left',vertical='center')
   if name=='当前 Full':c.fill=PatternFill('solid',fgColor='E2F0E7');c.font=Font(name='微软雅黑',bold=True,color='174B35')
   else:c.fill=PatternFill('solid',fgColor='EEF3F9' if j%2 else 'FFFFFF');c.font=Font(name='微软雅黑',color='18395E')
  if name=='ROOPF 原版':
   verify=[comparison[group,f,'original'] for f in range(1,n+1)]
   assert (wr,tr,lr)==(sum(float(x['mean_difference'])<0 for x in verify),sum(float(x['mean_difference'])==0 for x in verify),sum(float(x['mean_difference'])>0 for x in verify))
 abst=w[f'消融_{group}'];abst['A2']='移除 Anchor：移除预训练 Anchor 生成网络，两个参考位置以原生信赖域候选补位，温度控制失效。\n移除在线部分：移除六个在线生成算子和额外原生提议，保留代理、Residual 与 Agent。\n移除 Agent：移除整个 Agent 框架；其余两项沿用 Full 案例及准入规则，未为消融重建。'
 abst.row_dimensions[2].height=72
 header=next(i for i in range(1,abst.max_row+1) if abst.cell(i,1).value=='比较')
 for index,(label,arm) in enumerate([('Full vs 移除 Anchor','llm_no_anchor'),('Full vs 移除在线部分','llm_no_online_generators'),('Full vs 移除 Agent','no_agent')],header+1):
  c=summ['BBOB24' if group=='BBOB' else group,arm];values=[label,*[int(c[k]) for k in ['function_wins','function_ties','function_losses','trajectory_wins','trajectory_ties','trajectory_losses']]]
  for j,v in enumerate(values,1):abst.cell(index,j).value=v
 validation[group]={'functions':n,'algorithms':len(data),'rankings':[{ 'algorithm':a,'average_rank':scores[a]} for a in order],'ablation':[summ['BBOB24' if group=='BBOB' else group,a] for a in ['llm_no_anchor','llm_no_online_generators','no_agent']]}
notes=w['说明与来源'];updates={'当前版本':'本轮配对 Full：在线候选210 FE准入，LLM 210/230 FE调用；融合8轮，生成控制从下一轮起7轮，使用修复后输出契约和S01_L08匹配案例。','问题分类':'BBOB完整24函数（包含BBOB20），CEC-style 6函数，UAV 5函数。','汇总规则':'按问题类别分别统计，不新增跨类别合计表。','论文比较限制':'其他算法沿用论文归档数据，未在本轮重跑；任务偏移一致性仍未完全确认，排名仅供参考，不代表统计显著。','CEC-style 限制':'CEC-style六函数采用解锁执行流程；论文归档偏移一致性仍未完全确认。','UAV 限制':'UAV为项目恢复的五个路径规划场景；论文对照沿用已有归档数据。','在线消融限制':'移除全部六个在线生成算子及额外原生提议，保留代理、Residual、Agent与Full案例准入；案例未为消融重建。','原版后端说明':'本轮Full和三项消融使用相同CUDA代理后端；独立ROOPF原版基线保留CPU路径，后端差异包含在比较内。','实验协议':'35函数、每函数100条配对轨迹、D10、初始100点、总预算300 FE；三项保留Agent分支使用真实API。','Full 数据来源':'experiments/r19_s01l08_component_ablations/FUNCTION_STATISTICS.csv，当前配对基线来自本轮REFERENCES.npz。','消融汇总来源':'experiments/r19_s01l08_component_ablations/CATEGORY_SUMMARY.csv','逐函数消融来源':'experiments/r19_s01l08_component_ablations/FUNCTION_STATISTICS.csv、PAIRED_COMPARISONS.csv','源数据校验':'R19独立核验通过14000轨迹、2100请求和105个有效API覆盖组合；重新计算各问题组排名与胜平负。','记录日期':'2026-10-30','BBOB 实现说明':'完整24函数包含BBOB20；BBOB20具有批次耦合随机语义，34函数主分析需单独解释。'}
for row in notes.iter_rows(min_row=4):
 if row[0].value in updates:row[1].value=updates[row[0].value];row[1].alignment=Alignment(wrap_text=True,vertical='center')
notes.row_dimensions[4].height=50
w.save(TARGET)
check=load_workbook(TARGET,data_only=True)
for group,n in [('BBOB',24),('CEC-style',6),('UAV',5)]:
 s=check[f'均值方差_消融_{group}']; seen=defaultdict(int)
 for name,func,mean,var in s.iter_rows(min_row=4,values_only=True):
  if name:
   fid=int(func.replace(group,''));d=stats[group,fid,labels[name]]
   assert math.isclose(mean,float(d['mean']),rel_tol=1e-14,abs_tol=0) and math.isclose(var,float(d['variance']),rel_tol=1e-14,abs_tol=0);seen[name]+=1
 assert all(seen[k]==n for k in labels)
 s=check[f'对照_{group}'];assert [s.cell(i,3).value for i in range(7,s.max_row+1)]==sorted(s.cell(i,3).value for i in range(7,s.max_row+1))
 s=check[f'消融_{group}'];head=next(i for i in range(1,s.max_row+1) if s.cell(i,1).value=='比较')
 for offset,arm in enumerate(['llm_no_anchor','llm_no_online_generators','no_agent'],1):
  d=summ['BBOB24' if group=='BBOB' else group,arm];assert [s.cell(head+offset,j).value for j in range(2,8)]==[int(d[k]) for k in ['function_wins','function_ties','function_losses','trajectory_wins','trajectory_ties','trajectory_losses']]
record={'status':'verified','source_workbook_sha256':hashlib.sha256(SRC.read_bytes()).hexdigest(),'r19_independent_audit_sha256':hashlib.sha256((EXP/'INDEPENDENT_AUDIT.json').read_bytes()).hexdigest(),'r19_function_statistics_sha256':hashlib.sha256((EXP/'FUNCTION_STATISTICS.csv').read_bytes()).hexdigest(),'output_sha256':hashlib.sha256(TARGET.read_bytes()).hexdigest(),'details':validation,'note':'BBOB20 included in BBOB24 rankings and ablations; separate stochastic caveat retained. Paper baseline archived and CEC offsets unverified.'}
(OUT/'核验与来源.json').write_text(json.dumps(record,ensure_ascii=False,indent=2)+'\n')
print('RESULT',TARGET)
for g,d in validation.items():print(g,'top5',[(x['algorithm'],round(x['average_rank'],4)) for x in d['rankings'][:5]])
